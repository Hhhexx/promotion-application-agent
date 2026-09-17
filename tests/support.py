"""测试支撑工具（纯函数，不依赖 pytest）。

集中放三类东西，避免在多个测试文件里复制粘贴：
  1. 账本文本工厂（按模板生成抖加 / 垫付申请文本，供边界与异常路径造数）
  2. 账号表读取 / 比对工具（逐行逐列证明「写入未越界」）
  3. 端到端裁决流程的复用步骤
"""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

from app.repositories.xlsx_repository import SheetRow, TableSnapshot

#: 本批次「授权可写」的列（中文表头）。
#: 逐格比对时排除这些列——其余列必须与源表完全一致，证明写入没越界。
WRITABLE_HEADERS: frozenset[str] = frozenset(
    {"抖加", "抖加支付人", "打款人", "打款日期", "备注"}
)

#: 表头中文名 -> 仓储层英文键（与 app/repositories/sheet_schema.py 对齐）。
HEADER_TO_KEY: dict[str, str] = {
    "类型": "type",
    "抖音昵称": "douyin_nickname",
    "抖音号": "douyin_id",
    "报价": "quote_price",
    "抖加": "doujia_amount",
    "抖加支付人": "doujia_payer",
    "是否打款": "is_paid",
    "打款人": "payer",
    "打款日期": "pay_date",
    "发布日期": "publish_date",
    "是否接单": "accepted",
    "审核结果": "review_result",
    "备注": "remark",
}


# --------------------------------------------------------------------------- #
# 文件与单元格工具
# --------------------------------------------------------------------------- #
def sha256(path: str | Path) -> str:
    """文件字节级指纹 —— 用于「源表零污染」断言。"""
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def normalize_cell(value: Any) -> str:
    """把单元格值归一为可比较字符串（1 与 1.0、None 与 "" 视为相同）。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_sheet_rows(path: str | Path) -> dict[str, dict[str, Any]]:
    """直读 xlsx 第一个工作表，返回 {物理行号: {中文表头: 原值}}。

    整行全空的行（如台账第 12 行占位空行）按约定不参与比对，故跳过。
    """
    from openpyxl import load_workbook

    wb = load_workbook(Path(path), data_only=True)
    ws = wb.worksheets[0]
    headers = [c.value for c in ws[1]]
    out: dict[str, dict[str, Any]] = {}
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        out[str(idx)] = {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
    wb.close()
    return out


def make_snapshot(
    rows: list[SheetRow],
    *,
    empty_row_refs: list[str] | None = None,
    date_styles: dict[str, int] | None = None,
    highlighted_empty_remarks: list[str] | None = None,
) -> TableSnapshot:
    """构造内存快照，供校验层单测（不读盘）。"""
    return TableSnapshot(
        rows=rows,
        empty_row_refs=list(empty_row_refs or []),
        date_styles=dict(date_styles or {}),
        highlighted_empty_remarks=list(highlighted_empty_remarks or []),
        columns=[],
    )


# --------------------------------------------------------------------------- #
# 文本工厂（严格按 app/parsers 的正则模板生成）
# --------------------------------------------------------------------------- #
def make_doujia_text(
    items: list[tuple[str, object, int | None]],
    *,
    total: object,
    declared_count: int | None = None,
    declared_merged: int | None = None,
    date_str: str = "2026-09-02",
    payer: str = "爆火音乐",
) -> str:
    """生成抖加申请文本。`items` 元素为 `(博主名, 金额, 分笔数或 None)`。"""
    lines = [f"【抖加金额统计】今日已支付：¥{total}（{date_str}）", "", "明细如下："]
    for name, amount, split in items:
        seg = f"歌名：《你是我的仰望》- 博主名：{name} - 金额：¥{amount}"
        if split and split > 1:
            seg += f"（分 {split} 笔）"
        lines.append(seg)
    lines.append("")
    if declared_count is not None:
        tail = f"今日共 {declared_count} 笔支付申请"
        tail += f"，合并同博主后为 {declared_merged} 笔。" if declared_merged is not None else "。"
        lines.append(tail)
    lines.append(f"抖加支付人：{payer}")
    return "\n".join(lines) + "\n"


def make_advance_text(
    items: list[tuple[str, object]],
    *,
    total: object,
    declared_count: int | None = None,
    payer: str = "林老师",
    expected_pay_date: str | None = None,
    date_str: str = "2026-09-03",
) -> str:
    """生成垫付申请文本。`items` 元素为 `(博主名, 金额)`。"""
    lines = [f"【垫付统计】今日已垫付：¥{total}（{date_str}）", "", "明细如下："]
    for name, amount in items:
        lines.append(f"申请为【达人：{name}】垫付 ¥{amount}（歌曲《你是我的仰望》）")
    lines.append("")
    if declared_count is not None:
        lines.append(f"今日共 {declared_count} 笔垫付申请。")
    lines.append(f"打款人：{payer}")
    if expected_pay_date is not None:
        lines.append(f"预计打款日期：{expected_pay_date}")
    return "\n".join(lines) + "\n"


# --------------------------------------------------------------------------- #
# 复用流程
# --------------------------------------------------------------------------- #
def anomaly_by_code(payload: dict, code: str) -> dict:
    """从解析响应体里取指定异常码的第一条（不存在则断言失败）。"""
    for anomaly in payload.get("anomalies") or []:
        if anomaly.get("code") == code:
            return anomaly
    raise AssertionError(f"响应中不存在异常码 {code}：{payload.get('anomalies')}")


def p0_of(payload: dict) -> list[dict]:
    return [a for a in (payload.get("anomalies") or []) if a.get("severity") == "P0"]


def approve_advance(client, adv_data: dict) -> tuple[str, str]:
    """按 e2e 的顺序裁掉垫付批次的 3 个 P0，返回 (request_id, 被钉死的目标行号)。"""
    request_id = adv_data["request_id"]
    endpoints = f"/api/v1/requests/{request_id}/resolutions"

    client.post(
        endpoints,
        json={
            "anomaly_id": anomaly_by_code(adv_data, "COUNT_MISMATCH")["anomaly_id"],
            "decision": "accept",
            "operator": "qa",
            "note": "以明细为准",
        },
    )
    client.post(
        endpoints,
        json={
            "anomaly_id": anomaly_by_code(adv_data, "DATE_ORDER_INVALID")["anomaly_id"],
            "decision": "accept",
            "operator": "qa",
            "note": "确认为录入笔误",
        },
    )
    item = next(
        i for i in adv_data["line_items"] if (i.get("match_result") or {}).get("level") == "AMBIGUOUS"
    )
    pick = str(item["match_result"]["candidates"][0]["row_ref"])
    client.post(
        endpoints,
        json={
            "anomaly_id": anomaly_by_code(adv_data, "MATCH_AMBIGUOUS")["anomaly_id"],
            "decision": "override",
            "override_value": pick,
            "operator": "qa",
            "note": f"人工核对后指定第 {pick} 行",
        },
    )
    after = client.get(f"/api/v1/requests/{request_id}").json()["data"]
    for anomaly in after["anomalies"]:
        if anomaly.get("severity") == "P0" and not anomaly.get("resolved"):
            client.post(
                endpoints,
                json={
                    "anomaly_id": anomaly["anomaly_id"],
                    "decision": "accept",
                    "operator": "qa",
                    "note": "台账已有记录，以台账为准",
                },
            )
    return request_id, pick
