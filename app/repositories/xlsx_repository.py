"""账号表 xlsx 仓储（架构 §3 / §5.1 仓储层）。

**源表只读**：所有写入只作用于 `data/test-copies/` 下的副本；
仓储层不做业务校验（架构 §5.2 硬规则 2）。
"""

from __future__ import annotations

import shutil
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.domain.errors import ValidationError
from app.parsers.normalizer import classify_date_format, normalize_text
from app.repositories.sheet_schema import (
    KEY_BY_CN,
    LEDGER_HEADERS,
    LEDGER_SHEET,
    SHEET_COLUMNS,
    TRACE_COLUMNS,
    WritePolicy,
    assert_writable,
)

_DEFAULT_FILLS = (None, "", "00000000", "0", "FFFFFFFF")

#: 参与「日期写法统一性」体检的列。
DATE_COLUMNS: tuple[str, ...] = ("pay_date", "publish_date")


@dataclass
class SheetRow:
    """一行账号表数据（`row_ref` 为 Excel 物理行号，界面据此定位）。"""

    row_ref: str
    values: dict[str, Any] = field(default_factory=dict)

    def text(self, key: str) -> str:
        return normalize_text(self.values.get(key) or "")

    def is_blank(self, key: str) -> bool:
        return self.text(key) == ""


@dataclass
class TableSnapshot:
    """账号表一次性快照（读一次，供校验层与界面层共用，避免重复读盘）。"""

    rows: list[SheetRow] = field(default_factory=list)
    empty_row_refs: list[str] = field(default_factory=list)
    date_styles: dict[str, int] = field(default_factory=dict)
    highlighted_empty_remarks: list[str] = field(default_factory=list)
    columns: list[str] = field(default_factory=list)


def _is_highlighted(cell) -> bool:
    """单元格是否被底色高亮（用于识别「被标记但内容为空」的备注）。"""
    fill = getattr(cell, "fill", None)
    if fill is None or not getattr(fill, "fill_type", None):
        return False
    rgb = getattr(fill.start_color, "rgb", None)
    return bool(rgb) and str(rgb).upper() not in _DEFAULT_FILLS


class XlsxRepository:
    """一个 xlsx 文件（源表或副本）的读写封装。"""

    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)
        self.wb = load_workbook(self.path)
        self.ws: Worksheet = self.wb.worksheets[0]
        self.header: dict[str, int] = {}
        self._header_scan()
        self.highlighted_empty_remarks: list[str] = []

    # ------------------------------------------------------------------ 读
    def _header_scan(self) -> None:
        header: dict[str, int] = {}
        for col in range(1, self.ws.max_column + 1):
            value = self.ws.cell(row=1, column=col).value
            cn = normalize_text(str(value)) if value is not None else ""
            if cn in KEY_BY_CN:
                header[KEY_BY_CN[cn]] = col
        missing = [cn for cn, _key, _p in SHEET_COLUMNS if _key not in header]
        if missing:
            raise ValidationError(f"账号表缺少必需列：{'、'.join(missing)}")
        self.header = header

    def read_rows(self) -> tuple[list[SheetRow], list[str]]:
        """返回 (非空行, 整行空行的 row_ref 列表)。空行一律跳过不参与匹配。"""
        rows: list[SheetRow] = []
        empty_refs: list[str] = []
        self.highlighted_empty_remarks = []
        for r in range(2, self.ws.max_row + 1):
            values: dict[str, Any] = {}
            blank = True
            for key, col in self.header.items():
                cell = self.ws.cell(row=r, column=col)
                values[key] = cell.value
                if cell.value is not None and normalize_text(str(cell.value)) != "":
                    blank = False
            if blank:
                empty_refs.append(str(r))
                continue
            remark_cell = self.ws.cell(row=r, column=self.header["remark"])
            if (remark_cell.value in (None, "")) and _is_highlighted(remark_cell):
                self.highlighted_empty_remarks.append(str(r))
            rows.append(SheetRow(row_ref=str(r), values=values))
        return rows, empty_refs

    def date_format_styles(self) -> dict[str, int]:
        """统计**全部日期列**（打款日期 / 发布日期）的书写风格分布。

        `DATE_FORMAT_INCONSISTENT` 的口径是「台账内日期写法不统一」，因此跨列统计
        —— 只看单列会漏掉 `2026/8/21` 与 `4.10` 并存的真实情况。
        """
        styles: dict[str, int] = {}
        for key in DATE_COLUMNS:
            col = self.header.get(key)
            if col is None:
                continue
            for r in range(2, self.ws.max_row + 1):
                value = self.ws.cell(row=r, column=col).value
                if value is None or normalize_text(str(value)) == "":
                    continue
                style = classify_date_format(str(value))
                styles[style] = styles.get(style, 0) + 1
        return styles

    def columns_cn(self) -> list[str]:
        return [cn for cn, _k, _p in SHEET_COLUMNS]

    def snapshot(self) -> TableSnapshot:
        """产出一次性快照：行、空行引用、日期风格分布、高亮空备注。"""
        rows, empty_refs = self.read_rows()
        return TableSnapshot(
            rows=rows,
            empty_row_refs=empty_refs,
            date_styles=self.date_format_styles(),
            highlighted_empty_remarks=list(self.highlighted_empty_remarks),
            columns=self.columns_cn(),
        )

    # ------------------------------------------------------------------ 写
    def ensure_trace_columns(self) -> None:
        """补齐下划线前缀溯源列（业务 13 列之后追加，不改动业务列位置）。"""
        for name in TRACE_COLUMNS:
            if name in self.header:
                continue
            col = self.ws.max_column + 1
            self.ws.cell(row=1, column=col).value = name
            self.header[name] = col

    def write_row(
        self,
        row_ref: str,
        writes: dict[str, Any],
        trace: dict[str, str],
        remark: str | None = None,
    ) -> tuple[dict[str, Any], list[tuple[str, str]]]:
        """写一行。

        返回 `(实际写入, 被拦下的字段)`；被拦下者由编排层转为
        `TARGET_CELL_NONEMPTY`（P1，不覆盖已知值）。
        """
        row = int(row_ref)
        written: dict[str, Any] = {}
        blocked: list[tuple[str, str]] = []
        for key, value in writes.items():
            assert_writable(key)
            cell = self.ws.cell(row=row, column=self.header[key])
            existing = cell.value
            has_value = existing is not None and normalize_text(str(existing)) != ""
            if has_value and normalize_text(str(existing)) != normalize_text(str(value)):
                blocked.append((key, str(existing)))
                continue
            if not has_value:
                cell.value = value
            written[key] = cell.value

        if remark:
            cell = self.ws.cell(row=row, column=self.header["remark"])
            existing = normalize_text(str(cell.value)) if cell.value is not None else ""
            cell.value = f"{existing}；{remark}" if existing else remark
            written["remark"] = cell.value

        self.ensure_trace_columns()
        for key, value in trace.items():
            self.ws.cell(row=row, column=self.header[key]).value = value
        return written, blocked

    # --------------------------------------------------------------- 台账
    def ensure_ledger(self) -> Worksheet:
        """副本内自带写入台账，使幂等判重可跨进程重启（架构 §2.6）。"""
        if LEDGER_SHEET in self.wb.sheetnames:
            return self.wb[LEDGER_SHEET]
        sheet = self.wb.create_sheet(LEDGER_SHEET)
        sheet.append(LEDGER_HEADERS)
        sheet.sheet_state = "hidden"
        return sheet

    def ledger_keys(self) -> set[str]:
        if LEDGER_SHEET not in self.wb.sheetnames:
            return set()
        sheet = self.wb[LEDGER_SHEET]
        return {
            str(sheet.cell(row=r, column=1).value)
            for r in range(2, sheet.max_row + 1)
            if sheet.cell(row=r, column=1).value
        }

    def append_ledger(self, entries: list[dict[str, str]]) -> None:
        sheet = self.ensure_ledger()
        for entry in entries:
            sheet.append([entry.get(h, "") for h in LEDGER_HEADERS])

    def save(self, path: str | Path | None = None) -> Path:
        target = Path(path) if path else self.path
        self.wb.save(target)
        return target


def copy_test_copy(source: str | Path, dest: str | Path) -> Path:
    """复制源表为测试副本。`copy2` 保留原文件属性；此后只操作副本。"""
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, dest)
    return dest


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def policy_of(key: str) -> WritePolicy:
    return next(p for _cn, k, p in SHEET_COLUMNS if k == key)
