"""端到端验证的共享基座：断言登记、路径常量、源表直读。

本模块只提供**验证工具**，不含任何被测业务逻辑。
配套：`e2e_checks_intake.py`（链路进：解析/校验/闸门/裁决）、
`e2e_checks_output.py`（链路出：写表/汇总/汇报/审计），驱动入口 `e2e_smoke.py`。
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from pathlib import Path

#: 仓库根（本文件位于 tools/）。
ROOT = Path(__file__).resolve().parents[1]

#: 只读源表。整个验证过程中它的 sha256 必须保持不变。
SOURCE = ROOT / "data" / "你是我的仰望-推广账号表.xlsx"

#: 题目物料原文（与 data/samples/ 下的两份文本一致）。
DOUJIA = ROOT / "data" / "samples" / "抖加申请-2026-09-02.txt"
ADVANCE = ROOT / "data" / "samples" / "垫付申请-2026-09-03.txt"

#: 表头中文名 -> API 返回的行字段名（与 app/repositories/sheet_schema.py 对齐）。
HEADER_TO_KEY = {
    "类型": "type",
    "抖音昵称": "douyin_nickname",
    "抖音号": "douyin_id",
    "报价": "quote_price",
    "抖加": "doujia_amount",
    "抖加支付人": "doujia_payer",
    "是否打款": "is_paid",
    "打款人": "payer",
    "打款日期": "pay_date",
    "发布日期": "publish_date",
    "是否接单": "accepted",
    "审核结果": "review_result",
    "备注": "remark",
}

#: 两个批次**允许**写入的列。除此之外的列必须与源表逐行逐列一致。
WRITABLE_HEADERS = {"抖加", "抖加支付人", "打款人", "打款日期", "备注"}


# --------------------------------------------------------------------------- #
# 断言登记
# --------------------------------------------------------------------------- #
@dataclass
class Checker:
    """极简断言登记器：逐条打印，末尾汇总，失败即让进程退出码非 0。"""

    checks: int = 0
    failures: list[str] = field(default_factory=list)

    def check(self, label: str, condition: bool, detail: str = "") -> bool:
        self.checks += 1
        print(f"  [{'PASS' if condition else 'FAIL'}] {label}" + (f" -> {detail}" if detail else ""))
        if not condition:
            self.failures.append(label)
        return bool(condition)

    def title(self, text: str) -> None:
        print(f"\n=== {text} ===")

    def report(self) -> int:
        print("\n" + "=" * 68)
        if self.failures:
            print(f"结果：{self.checks - len(self.failures)}/{self.checks} 通过，{len(self.failures)} 项失败")
            for item in self.failures:
                print(f"  - {item}")
            return 1
        print(f"结果：{self.checks}/{self.checks} 全部通过")
        return 0


@dataclass
class Ctx:
    """跨两个 suite 传递的验证上下文（由 e2e_smoke.py 组装）。"""

    source_hash_before: str = ""
    table_id: str = ""
    copy_path: Path = Path()
    dj_id: str = ""
    adv_id: str = ""
    adv_data: dict = field(default_factory=dict)
    pick: str = ""
    adv_commit: dict = field(default_factory=dict)


# --------------------------------------------------------------------------- #
# 纯工具
# --------------------------------------------------------------------------- #
def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def norm_cell(value) -> str:
    """把 Excel 单元格值归一为可比较字符串（1 vs 1.0、None vs "" 视为相同）。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_source_rows(path: Path) -> dict[str, dict]:
    """用 openpyxl 直读源表，返回 {行号: {表头: 原始值}}。不走被测代码，避免自证。"""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    out: dict[str, dict] = {}
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        out[str(idx)] = {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
    wb.close()
    return out


# --------------------------------------------------------------------------- #
# 响应体取值
# --------------------------------------------------------------------------- #
def data_of(payload: dict) -> dict:
    return payload.get("data") or {}


def anomalies_of(payload: dict) -> list[dict]:
    """兼容两种入参：完整响应体 `{code,data,message}` 或直接的数据体 `{anomalies:[...]}`。"""
    if "anomalies" in payload:
        return payload.get("anomalies") or []
    return data_of(payload).get("anomalies") or []


def p0_of(payload: dict) -> list[dict]:
    return [a for a in anomalies_of(payload) if a.get("severity") == "P0"]


def by_code(payload: dict, code: str) -> list[dict]:
    return [a for a in anomalies_of(payload) if a.get("code") == code]


def one(items: list[dict]) -> dict:
    return items[0] if items else {}
